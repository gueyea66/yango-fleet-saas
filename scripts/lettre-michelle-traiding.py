# -*- coding: utf-8 -*-
"""Génère la lettre d'agrément Michelle Traiding en .docx et en .html (→ PDF)."""
import html, subprocess, pathlib

EXP = ["MICHELLE TRAIDING",
       "Sicap Liberté 3, n° 36/A — Dakar, Sénégal",
       "Tél. : +221 77 285 88 29",
       "RCCM SN DKR 2026 A 3767  ·  NINEA 012793768"]
DATE = "Dakar, le 29 juillet 2026"
DEST = ["FONSTAB", "À l'attention de Madame l'Administratrice",
        "Sphères ministérielles Ousmane Tanor Dieng", "Bâtiment C — Diamniadio, Sénégal"]
OBJET = "Objet : demande d'agrément — intendance et fournitures des manifestations"

CORPS = [
  "Madame l'Administratrice,",
  "Michelle Traiding intervient sur l'intendance des manifestations professionnelles. Notre rôle "
  "tient en une phrase : fournir, à la date fixée et au lieu convenu, tout ce dont la rencontre a "
  "besoin pour se tenir.",
]
POINTS = [
  ("1.", "Fournitures et dotations",
   "kits des participants, blocs et stylos, clés USB, supports imprimés, banderoles et "
   "kakémonos, cadeaux de protocole. Achat, personnalisation et livraison sur place."),
  ("2.", "Matériel de réception",
   "tentes, tables et chaises, sonorisation, groupes électrogènes, mobilier d'exposition : "
   "location, pose et reprise."),
  ("3.", "Véhicules",
   "voitures et bus avec chauffeur, pour les rotations pendant la rencontre comme pour les "
   "missions de terrain qui la prolongent."),
  ("4.", "Appui sur place",
   "équipe d'installation, personnel d'accueil et agents de liaison pendant toute la durée de "
   "la manifestation."),
]
FIN = [
  "Nous travaillons sur bon de commande : prix unitaire annoncé avant l'achat, facture "
  "accompagnée des pièces justificatives. Les délais que nous annonçons tiennent compte de ceux "
  "de nos fournisseurs, à Dakar comme à l'import.",
  "Nos activités déclarées couvrent le commerce général, l'import-export, la prestation de "
  "services, le transport routier ainsi que la vente et la location de véhicules : l'ensemble "
  "des prestations ci-dessus s'y rattache.",
  "Nous sollicitons votre agrément à ce titre et nous tenons à votre disposition pour tout "
  "complément.",
  "Pièces jointes : accusé d'immatriculation au RCCM, avis NINEA.",
  "Veuillez agréer, Madame l'Administratrice, l'expression de notre considération distinguée.",
]
SIGN = ["Mamadou Siradio DIALLO", "Gérant"]

BASE = pathlib.Path(__file__).resolve().parents[1] / "docs"
NOM = "lettre-agrement-fonstab-michelle-traiding"

# ---------------------------------------------------------------- HTML → PDF
points_html = "".join(
    f'    <div class="pt"><div class="n">{html.escape(n)}</div>'
    f'<div class="t"><b>{html.escape(t)}</b> — {html.escape(c)}</div></div>\n'
    for n, t, c in POINTS)

doc = f"""<!DOCTYPE html>
<html lang="fr">
<head>
<meta charset="utf-8" />
<title>Michelle Traiding — Demande d'agrément</title>
<style>
  @page {{ size: A4; margin: 18mm 24mm 16mm 24mm; }}
  body {{ font-family: Arial, Helvetica, sans-serif; font-size: 10.6pt; line-height: 1.42;
         color: #000; margin: 0; }}
  .exp {{ margin-bottom: 10mm; }}
  .exp .nom {{ font-weight: bold; letter-spacing: .05em; font-size: 12pt; }}
  .exp .ids {{ font-size: 8.4pt; color: #444; margin-top: 1.2mm; }}
  .date {{ text-align: right; margin-bottom: 7mm; }}
  .dest {{ margin-bottom: 7mm; }}
  .dest .nom {{ font-weight: bold; }}
  .objet {{ margin-bottom: 6mm; font-weight: bold; }}
  p {{ margin: 0 0 3.4mm; text-align: justify; }}
  .pts {{ margin: 1mm 0 4mm; }}
  .pt {{ display: flex; gap: 4mm; margin-bottom: 2.4mm; }}
  .pt .n {{ flex: 0 0 auto; width: 6mm; font-weight: bold; }}
  .pt .t {{ flex: 1 1 auto; text-align: justify; }}
  .sign {{ margin-top: 8mm; }}
  .sign .nom {{ font-weight: bold; margin-top: 13mm; }}
</style>
</head>
<body>
  <div class="exp">
    <div class="nom">{html.escape(EXP[0])}</div>
    {html.escape(EXP[1])}<br />
    {html.escape(EXP[2])}
    <div class="ids">{html.escape(EXP[3])}</div>
  </div>

  <div class="date">{html.escape(DATE)}</div>

  <div class="dest">
    <span class="nom">{html.escape(DEST[0])}</span><br />
    {html.escape(DEST[1])}<br />
    {html.escape(DEST[2])}<br />
    {html.escape(DEST[3])}
  </div>

  <div class="objet">{html.escape(OBJET)}</div>

""" + "".join(f"  <p>{html.escape(p)}</p>\n\n" for p in CORPS) \
    + f'  <div class="pts">\n{points_html}  </div>\n\n' \
    + "".join(f"  <p>{html.escape(p)}</p>\n\n" for p in FIN) + f"""  <div class="sign">
    <div class="nom">{html.escape(SIGN[0])}</div>
    <div>{html.escape(SIGN[1])}</div>
  </div>
</body>
</html>
"""
(BASE / f"{NOM}.html").write_text(doc, encoding="utf-8")

subprocess.run(["/opt/pw-browsers/chromium-1194/chrome-linux/chrome", "--headless", "--disable-gpu",
                "--no-sandbox", "--no-pdf-header-footer",
                f"--print-to-pdf={BASE / (NOM + '.pdf')}", f"file://{BASE / (NOM + '.html')}"],
               capture_output=True)

# ---------------------------------------------------------------- DOCX
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

d = Document()
sec = d.sections[0]
sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.6)
sec.left_margin = sec.right_margin = Cm(2.4)
st = d.styles["Normal"]
st.font.name, st.font.size = "Arial", Pt(10.5)
st.paragraph_format.space_after = Pt(9)
st.paragraph_format.line_spacing = 1.1
GRIS = RGBColor(0x44, 0x44, 0x44)


def para(txt="", bold=False, align=None, space_after=9, size=None, color=None):
    p = d.add_paragraph()
    if txt:
        r = p.add_run(txt); r.bold = bold
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    if align: p.alignment = align
    return p


para(EXP[0], bold=True, space_after=0, size=12)
para(EXP[1], space_after=0)
para(EXP[2], space_after=2)
para(EXP[3], space_after=24, size=8.5, color=GRIS)
para(DATE, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=18)
para(DEST[0], bold=True, space_after=0)
for l in DEST[1:]:
    para(l, space_after=0)
d.paragraphs[-1].paragraph_format.space_after = Pt(20)
para(OBJET, bold=True, space_after=16)
for p in CORPS:
    para(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
for n, t, c in POINTS:
    p = d.add_paragraph()
    r = p.add_run(f"{n} "); r.bold = True
    r = p.add_run(t); r.bold = True
    p.add_run(" — " + c)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Cm(0.5)
d.paragraphs[-1].paragraph_format.space_after = Pt(12)
for p in FIN:
    para(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
para("", space_after=26)
para(SIGN[0], bold=True, space_after=0)
para(SIGN[1], space_after=0)

d.save(BASE / f"{NOM}.docx")
print("ok")
