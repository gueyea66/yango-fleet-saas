# -*- coding: utf-8 -*-
"""Génère la lettre d'agrément Misaal Créa SARL en .docx et en .html (→ PDF)."""
import html, subprocess, pathlib

EXP = ["MISAAL CRÉA SARL",
       "Conception et organisation d'événements",
       "HLM 4, Villa n° 1385 — Dakar, Sénégal",
       "Tél. : +221 77 645 90 48",
       "RCCM SN DKR 2016 B 19405  ·  NINEA 006037908"]
DATE = "Dakar, le 12 juillet 2026"
DEST = ["FONSTAB", "Madame l'Administratrice",
        "Sphères ministérielles Ousmane Tanor Dieng", "Bâtiment C — Diamniadio, Sénégal"]
OBJET = "Objet : demande d'agrément — conception et organisation d'événements"

CORPS = [
  "Madame l'Administratrice,",
  "Nous sollicitons votre agrément en qualité d'organisateur d'événements.",
  "Misaal Créa imagine et réalise des manifestations professionnelles et institutionnelles. Notre "
  "point de départ n'est jamais un catalogue de prestations, mais une intention : ce que la "
  "rencontre doit produire, devant qui, dans quel délai et pour quel budget. Nous en tirons un "
  "concept, un déroulé et une scénographie, que nous soumettons avant tout engagement de dépense.",
  "Vient ensuite l'exécution, que nous assurons nous-mêmes : choix et réservation du lieu, "
  "identité visuelle de la rencontre — décor, habillage de scène, signalétique, supports "
  "imprimés —, constitution et suivi du fichier des invités, sélection et coordination des "
  "prestataires, direction de la séance et tenue du minutage, restauration, puis démontage et "
  "remise du site. L'équipe qui a conçu l'événement est celle qui le monte et qui le tient le "
  "jour dit.",
  "À la clôture, nous remettons un dossier complet : liste de présence, reportage photographique, "
  "film de restitution et récapitulatif des dépenses appuyé des justificatifs. Ces contenus sont "
  "produits par notre propre studio, sans intervenant extérieur, ce qui raccourcit les délais de "
  "livraison.",
  "Chaque projet est conduit par un chef de projet unique, joignable du premier rendez-vous "
  "jusqu'à la remise du dossier, sur un budget arrêté d'avance.",
  "Misaal Créa SARL exerce à Dakar depuis 2016. Ses activités déclarées portent sur la production "
  "vidéo — cinéma et télévision — ainsi que sur la formation. Les documents d'immatriculation "
  "sont joints à la présente.",
  "Nous nous tenons prêts à vous présenter nos réalisations et vous prions de croire, Madame "
  "l'Administratrice, à l'assurance de notre parfaite considération.",
]
SIGN = ["Le Gérant"]

BASE = pathlib.Path(__file__).resolve().parents[1] / "docs"
NOM = "lettre-agrement-fonstab-misaal-crea"

# ---------------------------------------------------------------- HTML → PDF
doc = f"""<!DOCTYPE html>
<html lang="fr">
<head>
<meta charset="utf-8" />
<title>Misaal Créa SARL — Demande d'agrément</title>
<style>
  @page {{ size: A4; margin: 16mm 23mm 14mm 23mm; }}
  body {{ font-family: "Helvetica Neue", Helvetica, Arial, sans-serif; font-size: 10.4pt;
         line-height: 1.45; color: #1a1a1a; margin: 0; }}
  .exp {{ text-align: right; margin-bottom: 10mm; }}
  .exp .nom {{ font-size: 13pt; font-weight: 600; letter-spacing: .06em; }}
  .exp .quoi {{ font-size: 8.6pt; color: #555; font-style: italic; margin-top: .8mm; }}
  .exp .adr {{ font-size: 8.6pt; color: #555; margin-top: 2.5mm; line-height: 1.55; }}
  .exp .ids {{ font-size: 8.2pt; color: #666; margin-top: 1mm; }}
  .dest {{ margin-bottom: 8mm; line-height: 1.45; }}
  .dest .nom {{ font-weight: 600; }}
  .date {{ margin-bottom: 7mm; }}
  .objet {{ margin-bottom: 6mm; font-weight: 600; }}
  p {{ margin: 0 0 3.4mm; text-align: justify; }}
  .sign {{ margin-top: 8mm; }}
  .sign .role {{ font-weight: 600; }}
</style>
</head>
<body>
  <div class="exp">
    <div class="nom">{html.escape(EXP[0])}</div>
    <div class="quoi">{html.escape(EXP[1])}</div>
    <div class="adr">{html.escape(EXP[2])}<br />{html.escape(EXP[3])}</div>
    <div class="ids">{html.escape(EXP[4])}</div>
  </div>

  <div class="dest">
    <span class="nom">{html.escape(DEST[0])}</span><br />
    {html.escape(DEST[1])}<br />
    {html.escape(DEST[2])}<br />
    {html.escape(DEST[3])}
  </div>

  <div class="date">{html.escape(DATE)}</div>

  <div class="objet">{html.escape(OBJET)}</div>

""" + "".join(f"  <p>{html.escape(p)}</p>\n\n" for p in CORPS) + f"""  <div class="sign">
    <div class="role">{html.escape(SIGN[0])}</div>
  </div>
</body>
</html>
"""
(BASE / f"{NOM}.html").write_text(doc, encoding="utf-8")
subprocess.run(["/opt/pw-browsers/chromium-1194/chrome-linux/chrome", "--headless", "--disable-gpu",
                "--no-sandbox", "--no-pdf-header-footer",
                f"--print-to-pdf={BASE / (NOM + '.pdf')}", f"file://{BASE / (NOM + '.html')}"],
               capture_output=True)

# ---------------------------------------------------------------- DOCX
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

d = Document()
sec = d.sections[0]
sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.4)
sec.left_margin = sec.right_margin = Cm(2.3)
st = d.styles["Normal"]
st.font.name, st.font.size = "Calibri", Pt(11)
st.paragraph_format.space_after = Pt(9)
st.paragraph_format.line_spacing = 1.12
GRIS = RGBColor(0x55, 0x55, 0x55)
R = WD_ALIGN_PARAGRAPH.RIGHT
J = WD_ALIGN_PARAGRAPH.JUSTIFY


def para(txt="", bold=False, italic=False, align=None, space_after=9, size=None, color=None):
    p = d.add_paragraph()
    if txt:
        r = p.add_run(txt); r.bold = bold; r.italic = italic
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    if align: p.alignment = align
    return p


para(EXP[0], bold=True, align=R, space_after=0, size=13)
para(EXP[1], italic=True, align=R, space_after=6, size=8.5, color=GRIS)
para(EXP[2], align=R, space_after=0, size=8.5, color=GRIS)
para(EXP[3], align=R, space_after=2, size=8.5, color=GRIS)
para(EXP[4], align=R, space_after=24, size=8, color=GRIS)
para(DEST[0], bold=True, space_after=0)
for l in DEST[1:]:
    para(l, space_after=0)
d.paragraphs[-1].paragraph_format.space_after = Pt(20)
para(DATE, space_after=16)
para(OBJET, bold=True, space_after=14)
for p in CORPS:
    para(p, align=J)
para("", space_after=22)
para(SIGN[0], bold=True, space_after=0)

d.save(BASE / f"{NOM}.docx")
print("ok")
