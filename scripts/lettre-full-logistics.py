# -*- coding: utf-8 -*-
"""Génère la lettre d'agrément Full Logistics SARL en .docx et en .html (→ PDF)."""
import html, subprocess, pathlib

EXP = ["FULL LOGISTICS SARL",
       "Société à responsabilité limitée au capital de 100 000 F CFA",
       "HLM Mariste, Îlot Z, Villa n° 4 — Dakar, Sénégal",
       "Tél. : +221 77 645 65 17",
       "RCCM SN DKR 2024 B 11455  ·  NINEA 011088152"]
DATE = "Dakar, le 29 juin 2026"
DEST = ["FONSTAB", "Madame l'Administratrice",
        "Sphères ministérielles Ousmane Tanor Dieng", "Bâtiment C — Diamniadio, Sénégal"]
OBJET = "Objet : demande d'agrément en qualité d'opérateur événementiel"

INTRO = [
  "Madame l'Administratrice,",
  "Nous avons l'honneur de solliciter votre agrément pour la conduite complète de vos "
  "manifestations.",
  "Full Logistics prend une rencontre à sa charge du cadrage jusqu'au dossier de clôture. Nous "
  "en arrêtons le budget et le calendrier, retenons le site, mobilisons les moyens et les "
  "équipes, tenons la séance, puis rendons les lieux et le compte de la dépense. Le "
  "commanditaire garde la décision ; nous portons l'exécution.",
]

BLOCS = [
  ("Cadrage.",
   "Étude du besoin, repérage et réservation du site, budget par poste, rétroplanning, "
   "préparation et suivi des invitations."),
  ("Montage.",
   "Transport et pose des équipements, agencement des salles et des espaces d'exposition, "
   "régie du son et de l'image, balisage des accès."),
  ("Conduite de la séance.",
   "Enregistrement des participants, tenue du programme, liaison avec le protocole, "
   "restauration, prise en charge des délégations et de leurs déplacements."),
  ("Clôture.",
   "Démontage, restitution du site dans son état initial, dossier de fin de mission : liste de "
   "présence, photographies, dépenses justifiées."),
]

FIN = [
  "Nos véhicules, nos équipes de manutention et notre réseau de prestataires nous permettent de "
  "tenir une manifestation à Dakar comme en région, y compris lorsqu'elle se déplace d'un site à "
  "l'autre. Un coordonnateur est désigné à la commande et reste votre interlocuteur jusqu'à la "
  "remise du dossier.",
  "Nos activités déclarées couvrent la manutention et le shipping, le transport routier, le "
  "transport de personnes et le tourisme, ainsi que le commerce général : les prestations "
  "ci-dessus s'y rattachent.",
  "Les documents d'immatriculation de la société sont joints à la présente.",
  "Dans l'attente d'une suite favorable, nous vous prions d'agréer, Madame l'Administratrice, "
  "l'expression de notre haute considération.",
]
SIGN = ["Babacar GUEYE", "Gérant"]

BASE = pathlib.Path(__file__).resolve().parents[1] / "docs"
NOM = "lettre-agrement-fonstab-full-logistics"

# ---------------------------------------------------------------- HTML → PDF
blocs_html = "".join(
    f"  <p><b>{html.escape(t)}</b> {html.escape(c)}</p>\n\n" for t, c in BLOCS)

doc = f"""<!DOCTYPE html>
<html lang="fr">
<head>
<meta charset="utf-8" />
<title>Full Logistics SARL — Demande d'agrément</title>
<style>
  @page {{ size: A4; margin: 15mm 23mm 12mm 23mm; }}
  body {{ font-family: Georgia, "Times New Roman", serif; font-size: 10.5pt; line-height: 1.36;
         color: #111; margin: 0; }}
  .exp {{ text-align: center; padding-bottom: 3mm; }}
  .exp .nom {{ font-size: 15pt; letter-spacing: .22em; font-variant: small-caps; font-weight: bold; }}
  .exp .legal {{ font-size: 8.2pt; color: #555; margin-top: 1.5mm; line-height: 1.5; }}
  .exp .ids {{ font-size: 8.2pt; color: #555; margin-top: 1mm; }}
  .double {{ border-top: 1.6pt solid #111; border-bottom: .6pt solid #111; height: 1.2mm;
             margin-bottom: 7mm; }}
  .date {{ text-align: right; margin-bottom: 6mm; }}
  .dest {{ margin-bottom: 6mm; line-height: 1.4; }}
  .dest .nom {{ font-weight: bold; }}
  .objet {{ margin-bottom: 6mm; }}
  .objet b {{ font-variant: small-caps; letter-spacing: .03em; }}
  p {{ margin: 0 0 3mm; text-align: justify; }}
  .sign {{ margin-top: 6mm; }}
  .sign .nom {{ font-weight: bold; margin-top: 11mm; }}
  .sign .role {{ font-size: 9.6pt; color: #555; }}
</style>
</head>
<body>
  <div class="exp">
    <div class="nom">{html.escape(EXP[0])}</div>
    <div class="legal">{html.escape(EXP[1])}<br />{html.escape(EXP[2])} &nbsp;—&nbsp; {html.escape(EXP[3])}</div>
    <div class="ids">{html.escape(EXP[4])}</div>
  </div>
  <div class="double"></div>

  <div class="date">{html.escape(DATE)}</div>

  <div class="dest">
    <span class="nom">{html.escape(DEST[0])}</span><br />
    {html.escape(DEST[1])}<br />
    {html.escape(DEST[2])}<br />
    {html.escape(DEST[3])}
  </div>

  <div class="objet"><b>{html.escape(OBJET)}</b></div>

""" + "".join(f"  <p>{html.escape(p)}</p>\n\n" for p in INTRO) + blocs_html \
    + "".join(f"  <p>{html.escape(p)}</p>\n\n" for p in FIN) + f"""  <div class="sign">
    <div class="nom">{html.escape(SIGN[0])}</div>
    <div class="role">{html.escape(SIGN[1])}</div>
  </div>
</body>
</html>
"""
(BASE / f"{NOM}.html").write_text(doc, encoding="utf-8")

subprocess.run(["/opt/pw-browsers/chromium-1194/chrome-linux/chrome", "--headless", "--disable-gpu",
                "--no-sandbox", "--no-pdf-header-footer",
                f"--print-to-pdf={BASE / (NOM + '.pdf')}", f"file://{BASE / (NOM + '.html')}"],
               capture_output=True)

# ---------------------------------------------------------------- DOCX
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

d = Document()
sec = d.sections[0]
sec.top_margin, sec.bottom_margin = Cm(1.7), Cm(1.5)
sec.left_margin = sec.right_margin = Cm(2.3)
st = d.styles["Normal"]
st.font.name, st.font.size = "Georgia", Pt(10.5)
st.paragraph_format.space_after = Pt(9)
st.paragraph_format.line_spacing = 1.1
GRIS = RGBColor(0x55, 0x55, 0x55)
C = WD_ALIGN_PARAGRAPH.CENTER
J = WD_ALIGN_PARAGRAPH.JUSTIFY


def para(txt="", bold=False, align=None, space_after=9, size=None, color=None):
    p = d.add_paragraph()
    if txt:
        r = p.add_run(txt); r.bold = bold
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    if align: p.alignment = align
    return p


para(EXP[0], bold=True, align=C, space_after=2, size=14)
para(EXP[1], align=C, space_after=0, size=8, color=GRIS)
para(f"{EXP[2]} — {EXP[3]}", align=C, space_after=0, size=8, color=GRIS)
para(EXP[4], align=C, space_after=22, size=8, color=GRIS)
para(DATE, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=16)
para(DEST[0], bold=True, space_after=0)
for l in DEST[1:]:
    para(l, space_after=0)
d.paragraphs[-1].paragraph_format.space_after = Pt(18)
para(OBJET, bold=True, space_after=14)
for p in INTRO:
    para(p, align=J)
for titre, corps in BLOCS:
    p = d.add_paragraph()
    r = p.add_run(titre + " "); r.bold = True
    p.add_run(corps)
    p.alignment = J
    p.paragraph_format.space_after = Pt(9)
for p in FIN:
    para(p, align=J)
para("", space_after=24)
para(SIGN[0], bold=True, space_after=0)
para(SIGN[1], space_after=0, size=9.5, color=GRIS)

d.save(BASE / f"{NOM}.docx")
print("ok")
